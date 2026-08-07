import csv
import hashlib
import io
import json
import re

import pdfplumber
import openpyxl
import xlrd
from docx import Document
from unidecode import unidecode

from application.database import db
from application.models.models import KnowledgeBase, KnowledgeChunk, KnowledgeDocument


class KnowledgeServiceError(Exception):
    def __init__(self, message, status_code=400):
        super().__init__(message)
        self.status_code = status_code


def normalize_search_text(value):
    value = unidecode(str(value or "")).lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9\s]", " ", value)).strip()


def chunk_text(value, size=1400, overlap=180):
    text = re.sub(r"\r\n?", "\n", str(value or "")).strip()
    if not text:
        return []

    size = max(int(size), 300)
    overlap = max(0, min(int(overlap), size // 3))
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
    chunks = []
    current = ""

    def append_piece(piece):
        nonlocal current
        if not current:
            current = piece
            return
        candidate = current + "\n\n" + piece
        if len(candidate) <= size:
            current = candidate
            return
        chunks.append(current.strip())
        prefix = current[-overlap:].lstrip() if overlap else ""
        current = (prefix + "\n\n" + piece).strip() if prefix else piece

    for paragraph in paragraphs:
        if len(paragraph) <= size:
            append_piece(paragraph)
            continue
        sentences = re.split(r"(?<=[.!?])\s+", paragraph)
        for sentence in sentences:
            sentence = sentence.strip()
            while len(sentence) > size:
                append_piece(sentence[:size])
                sentence = sentence[size - overlap:] if overlap else sentence[size:]
            if sentence:
                append_piece(sentence)

    if current:
        chunks.append(current.strip())
    return [item for item in chunks if item]


class KnowledgeService(object):
    SUPPORTED_EXTENSIONS = {
        ".pdf", ".txt", ".md", ".markdown", ".csv", ".json",
        ".docx", ".xlsx", ".xls",
    }

    def __init__(self, app):
        self.app = app

    @staticmethod
    def _serialize_base(item):
        return {
            "id": str(item.id),
            "code": item.code,
            "name": item.name,
            "description": item.description,
            "access_scope": item.access_scope,
            "allowed_department_ids": item.allowed_department_ids or [],
            "active": bool(item.active),
            "created_at": item.created_at,
            "updated_at": item.updated_at,
        }

    @staticmethod
    def _serialize_document(item, include_text=False):
        result = {
            "id": str(item.id),
            "knowledge_base_id": str(item.knowledge_base_id),
            "title": item.title,
            "file_name": item.file_name,
            "mime_type": item.mime_type,
            "source_type": item.source_type,
            "source_url": item.source_url,
            "checksum": item.checksum,
            "version": item.version,
            "status": item.status,
            "uploaded_by": item.uploaded_by,
            "properties": item.properties or {},
            "created_at": item.created_at,
            "updated_at": item.updated_at,
        }
        if include_text:
            result["extracted_text"] = item.extracted_text
        return result

    def get_base(self, tenant_id, base_id):
        item = KnowledgeBase.query.filter(
            KnowledgeBase.id == base_id,
            KnowledgeBase.tenant_id == tenant_id,
            KnowledgeBase.deleted.is_(False),
        ).first()
        if item is None:
            raise KnowledgeServiceError("Không tìm thấy kho tri thức.", 404)
        return item

    def list_bases(self, tenant_id):
        items = KnowledgeBase.query.filter(
            KnowledgeBase.tenant_id == tenant_id,
            KnowledgeBase.deleted.is_(False),
        ).order_by(KnowledgeBase.created_at.desc()).all()
        return [self._serialize_base(item) for item in items]

    def create_base(self, tenant_id, payload):
        code = str(payload.get("code") or "").strip()
        name = str(payload.get("name") or "").strip()
        if not code or not name:
            raise KnowledgeServiceError("Mã và tên kho tri thức là bắt buộc.")
        if KnowledgeBase.query.filter(
            KnowledgeBase.tenant_id == tenant_id,
            KnowledgeBase.code == code,
            KnowledgeBase.deleted.is_(False),
        ).first():
            raise KnowledgeServiceError("Mã kho tri thức đã tồn tại.", 409)
        item = KnowledgeBase(
            tenant_id=tenant_id,
            code=code,
            name=name,
            description=payload.get("description"),
            access_scope=str(payload.get("access_scope") or "COMPANY").upper(),
            allowed_department_ids=payload.get("allowed_department_ids") or [],
            active=payload.get("active", True) is not False,
            properties=payload.get("properties") or {},
        )
        db.session.add(item)
        db.session.commit()
        return self._serialize_base(item)

    def get_or_create_base(self, tenant_id, code, name, description="", properties=None):
        item = KnowledgeBase.query.filter(
            KnowledgeBase.tenant_id == tenant_id,
            KnowledgeBase.code == code,
            KnowledgeBase.deleted.is_(False),
        ).first()
        if item is not None:
            return item
        item = KnowledgeBase(
            tenant_id=tenant_id,
            code=code,
            name=name,
            description=description,
            access_scope="COMPANY",
            allowed_department_ids=[],
            active=True,
            properties=properties or {},
        )
        db.session.add(item)
        db.session.commit()
        return item

    def list_documents(self, tenant_id, knowledge_base_id=None):
        query = KnowledgeDocument.query.filter(
            KnowledgeDocument.tenant_id == tenant_id,
            KnowledgeDocument.deleted.is_(False),
        )
        if knowledge_base_id:
            query = query.filter(KnowledgeDocument.knowledge_base_id == knowledge_base_id)
        items = query.order_by(KnowledgeDocument.created_at.desc()).all()
        return [self._serialize_document(item) for item in items]

    def extract_file(self, file_name, mime_type, body):
        max_size = self.app.config.get("CHATBOT_MAX_KNOWLEDGE_FILE_SIZE", 20 * 1024 * 1024)
        if not body:
            raise KnowledgeServiceError("Tệp tải lên đang trống.")
        if len(body) > max_size:
            raise KnowledgeServiceError("Tệp vượt quá giới hạn {} MB.".format(max_size // 1024 // 1024), 413)

        lowered = str(file_name or "").lower()
        extension = "." + lowered.rsplit(".", 1)[-1] if "." in lowered else ""
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise KnowledgeServiceError("Chỉ hỗ trợ PDF, DOCX, XLS/XLSX, TXT, Markdown, CSV và JSON.", 415)

        if extension == ".pdf":
            try:
                with pdfplumber.open(io.BytesIO(body)) as document:
                    pages = [(page.extract_text() or "").strip() for page in document.pages]
                content = "\n\n".join(page for page in pages if page)
            except Exception as error:
                raise KnowledgeServiceError("Không thể đọc tệp PDF: {}".format(error), 422)
        elif extension == ".docx":
            try:
                document = Document(io.BytesIO(body))
                blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
                for table in document.tables:
                    for row in table.rows:
                        blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
                content = "\n".join(blocks)
            except Exception as error:
                raise KnowledgeServiceError("Không thể đọc tệp DOCX: {}".format(error), 422)
        elif extension == ".xlsx":
            try:
                workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
                rows = []
                for sheet in workbook.worksheets:
                    rows.append("[Sheet: {}]".format(sheet.title))
                    for row in sheet.iter_rows(values_only=True):
                        rows.append(" | ".join("" if value is None else str(value) for value in row))
                content = "\n".join(rows)
            except Exception as error:
                raise KnowledgeServiceError("Không thể đọc tệp XLSX: {}".format(error), 422)
        elif extension == ".xls":
            try:
                workbook = xlrd.open_workbook(file_contents=body)
                rows = []
                for sheet in workbook.sheets():
                    rows.append("[Sheet: {}]".format(sheet.name))
                    for row_index in range(sheet.nrows):
                        rows.append(" | ".join(str(value) for value in sheet.row_values(row_index)))
                content = "\n".join(rows)
            except Exception as error:
                raise KnowledgeServiceError("Không thể đọc tệp XLS: {}".format(error), 422)
        else:
            content = body.decode("utf-8-sig", errors="replace")
            if extension == ".json":
                try:
                    content = json.dumps(json.loads(content), ensure_ascii=False, indent=2)
                except ValueError:
                    raise KnowledgeServiceError("Tệp JSON không hợp lệ.", 422)
            elif extension == ".csv":
                try:
                    rows = list(csv.reader(io.StringIO(content)))
                    content = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
                except csv.Error as error:
                    raise KnowledgeServiceError("Tệp CSV không hợp lệ: {}".format(error), 422)

        if not content.strip():
            raise KnowledgeServiceError("Không trích xuất được nội dung chữ từ tệp.", 422)
        return content.strip(), mime_type or "application/octet-stream"

    def ingest_text(self, tenant_id, knowledge_base_id, title, content, **metadata):
        self.get_base(tenant_id, knowledge_base_id)
        title = str(title or "").strip()
        content = str(content or "").strip()
        if not title or not content:
            raise KnowledgeServiceError("Tiêu đề và nội dung tài liệu là bắt buộc.")

        checksum_source = content
        if str(metadata.get("source_type") or "").upper().startswith("CHAT_"):
            properties = metadata.get("properties") or {}
            checksum_source = "{}\n{}\n{}".format(
                properties.get("conversation_id") or "",
                properties.get("message_id") or "",
                content,
            )
        checksum = hashlib.sha256(checksum_source.encode("utf-8")).hexdigest()
        duplicate = KnowledgeDocument.query.filter(
            KnowledgeDocument.tenant_id == tenant_id,
            KnowledgeDocument.knowledge_base_id == knowledge_base_id,
            KnowledgeDocument.checksum == checksum,
            KnowledgeDocument.deleted.is_(False),
        ).first()
        if duplicate:
            result = self._serialize_document(duplicate)
            result["duplicate"] = True
            return result

        pieces = chunk_text(
            content,
            self.app.config.get("CHATBOT_CHUNK_SIZE", 1400),
            self.app.config.get("CHATBOT_CHUNK_OVERLAP", 180),
        )
        if not pieces:
            raise KnowledgeServiceError("Tài liệu không có nội dung có thể lập chỉ mục.", 422)

        item = KnowledgeDocument(
            tenant_id=tenant_id,
            knowledge_base_id=knowledge_base_id,
            title=title,
            file_name=metadata.get("file_name"),
            mime_type=metadata.get("mime_type"),
            source_type=metadata.get("source_type") or "TEXT",
            source_url=metadata.get("source_url"),
            checksum=checksum,
            status="READY",
            extracted_text=content,
            uploaded_by=metadata.get("uploaded_by"),
            properties=metadata.get("properties") or {},
        )
        try:
            db.session.add(item)
            db.session.flush()
            for index, piece in enumerate(pieces):
                db.session.add(KnowledgeChunk(
                    tenant_id=tenant_id,
                    knowledge_base_id=knowledge_base_id,
                    document_id=item.id,
                    chunk_index=index,
                    content=piece,
                    content_search=normalize_search_text(piece),
                    token_count=len(piece.split()),
                    properties={},
                ))
            db.session.commit()
        except Exception:
            db.session.rollback()
            raise
        result = self._serialize_document(item)
        result["chunk_count"] = len(pieces)
        return result

    def ingest_chat_message(self, tenant_id, conversation_id, message_id, sender_id, content, participant_ids=None):
        knowledge_base = self.get_or_create_base(
            tenant_id,
            "chat-derived",
            "Kiến thức từ hội thoại",
            "Tin nhắn và tệp chat được lập chỉ mục tự động theo quyền thành viên cuộc trò chuyện.",
            properties={"managed": True, "source": "chat"},
        )
        return self.ingest_text(
            tenant_id,
            knowledge_base.id,
            "Tin nhắn {}".format(message_id or "chat"),
            content,
            source_type="CHAT_MESSAGE",
            source_url="tinode://{}".format(conversation_id or ""),
            uploaded_by=str(sender_id or "chat-user"),
            properties={
                "conversation_id": str(conversation_id or ""),
                "message_id": str(message_id or ""),
                "sender_id": str(sender_id or ""),
                "allowed_user_ids": [str(value) for value in (participant_ids or []) if value],
            },
        )

    def delete_document(self, tenant_id, document_id):
        item = KnowledgeDocument.query.filter(
            KnowledgeDocument.id == document_id,
            KnowledgeDocument.tenant_id == tenant_id,
            KnowledgeDocument.deleted.is_(False),
        ).first()
        if item is None:
            raise KnowledgeServiceError("Không tìm thấy tài liệu.", 404)
        item.deleted = True
        for chunk in KnowledgeChunk.query.filter(
            KnowledgeChunk.document_id == item.id,
            KnowledgeChunk.deleted.is_(False),
        ).all():
            chunk.deleted = True
        db.session.commit()

    def retrieve(
        self,
        query_text,
        tenant_id,
        knowledge_base_id=None,
        department_id=None,
        user_ids=None,
        limit=None,
        exclude_source_prefixes=None,
    ):
        normalized_query = normalize_search_text(query_text)
        terms = [term for term in normalized_query.split() if len(term) > 1]
        if not terms:
            return []

        query = db.session.query(KnowledgeChunk, KnowledgeDocument, KnowledgeBase).join(
            KnowledgeDocument, KnowledgeChunk.document_id == KnowledgeDocument.id
        ).join(KnowledgeBase, KnowledgeChunk.knowledge_base_id == KnowledgeBase.id).filter(
            KnowledgeChunk.tenant_id == tenant_id,
            KnowledgeChunk.deleted.is_(False),
            KnowledgeDocument.deleted.is_(False),
            KnowledgeDocument.status == "READY",
            KnowledgeBase.deleted.is_(False),
            KnowledgeBase.active.is_(True),
        )
        if knowledge_base_id:
            query = query.filter(KnowledgeChunk.knowledge_base_id == knowledge_base_id)
        candidates = query.order_by(KnowledgeChunk.updated_at.desc()).limit(
            self.app.config.get("CHATBOT_RETRIEVAL_CANDIDATES", 500)
        ).all()

        results = []
        unique_terms = set(terms)
        viewer_ids = {str(value) for value in (user_ids or []) if value}
        excluded_prefixes = tuple(
            str(value).upper()
            for value in (exclude_source_prefixes or [])
            if str(value).strip()
        )
        for chunk, document, knowledge_base in candidates:
            source_type = str(document.source_type or "").upper()
            if excluded_prefixes and source_type.startswith(excluded_prefixes):
                continue
            scope = str(knowledge_base.access_scope or "COMPANY").upper()
            allowed_departments = [str(value) for value in (knowledge_base.allowed_department_ids or [])]
            if scope == "DEPARTMENT" and (not department_id or str(department_id) not in allowed_departments):
                continue
            document_properties = document.properties or {}
            allowed_users = {str(value) for value in (document_properties.get("allowed_user_ids") or []) if value}
            if source_type.startswith("CHAT_") and allowed_users and not viewer_ids.intersection(allowed_users):
                continue
            searchable = chunk.content_search or normalize_search_text(chunk.content)
            matched = [term for term in unique_terms if term in searchable]
            if not matched:
                continue
            coverage = len(matched) / max(len(unique_terms), 1)
            frequency = sum(searchable.count(term) for term in matched)
            phrase_bonus = 4.0 if normalized_query in searchable else 0.0
            score = round(coverage * 10.0 + min(frequency, 10) * 0.2 + phrase_bonus, 4)
            results.append({
                "chunk_id": str(chunk.id),
                "document_id": str(document.id),
                "knowledge_base_id": str(knowledge_base.id),
                "title": document.title,
                "file_name": document.file_name,
                "page_number": chunk.page_number,
                "content": chunk.content,
                "score": score,
            })
        results.sort(key=lambda item: item["score"], reverse=True)
        return results[:int(limit or self.app.config.get("CHATBOT_RETRIEVAL_LIMIT", 6))]

    @staticmethod
    def format_context(matches):
        blocks = []
        for index, item in enumerate(matches, 1):
            label = item.get("title") or item.get("file_name") or "Tài liệu"
            blocks.append("[Nguồn {}: {}]\n{}".format(index, label, item.get("content") or ""))
        return "\n\n".join(blocks)
